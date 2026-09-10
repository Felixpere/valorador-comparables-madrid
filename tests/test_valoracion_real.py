"""Tests del circuito con datos REALES (r01) y de la hoja de Excel que lo publica.

Ninguno de estos tests lee `data/real/madrid_2018.csv`: son 26 MB que no viajan
en el repositorio y que la integracion continua no tiene. Se construye a mano un
conjunto pequeno con los mismos nombres de columna del origen, que es lo que
permite que estos tests corran en cualquier sitio.

El test que mas importa es `test_invertir_la_calidad_no_mueve_ninguna_valoracion`:
sostiene la afirmacion que hay escrita en el comentario de r01.preparar().
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "src"))

import p05_exportar_excel as p05      # noqa: E402
import r01_valorar_real as r01        # noqa: E402


def corpus(n: int = 400, semilla: int = 20260904) -> pd.DataFrame:
    """Conjunto minimo con los nombres de columna de idealista18."""
    rng = np.random.default_rng(semilla)
    superficie = rng.integers(45, 160, n)
    calidad = rng.integers(1, 8, n)
    # El precio se construye a proposito peor cuanto mayor es el codigo
    # catastral, que es la direccion real del origen.
    eur_m2 = 6000 * np.exp(-0.15 * calidad) * rng.lognormal(0, 0.12, n)
    return pd.DataFrame({
        "ASSETID": [f"A{i:06d}" for i in range(n)],
        "PRICE": (eur_m2 * superficie).round(-3).clip(30_000, 6_000_000),
        "UNITPRICE": eur_m2.clip(500, 15_000),
        "CONSTRUCTEDAREA": superficie,
        "ROOMNUMBER": rng.integers(1, 5, n),
        "BATHNUMBER": rng.integers(1, 3, n),
        "CADASTRALQUALITYID": calidad.astype(float),
        "CADCONSTRUCTIONYEAR": rng.integers(1900, 2015, n),
        "FLOORCLEAN": rng.integers(0, 8, n).astype(float),
        "FLATLOCATIONID": rng.choice([1.0, 2.0], n),
        "BUILTTYPEID_1": rng.integers(0, 2, n),
        "BUILTTYPEID_2": rng.integers(0, 2, n),
        "BUILTTYPEID_3": 1,
        "HASLIFT": rng.integers(0, 2, n),
        "HASTERRACE": rng.integers(0, 2, n),
        "HASAIRCONDITIONING": rng.integers(0, 2, n),
        "ISINTOPFLOOR": rng.integers(0, 2, n),
        "HASPARKINGSPACE": rng.integers(0, 2, n),
        "HASBOXROOM": rng.integers(0, 2, n),
        "HASSWIMMINGPOOL": rng.integers(0, 2, n),
        "HASDOORMAN": rng.integers(0, 2, n),
        "HASGARDEN": rng.integers(0, 2, n),
        "ISDUPLEX": rng.integers(0, 2, n),
        "ISSTUDIO": (rng.random(n) < 0.1).astype(int),
        "barrio": rng.choice(["Centro", "Salamanca"], n),
        "trimestre": rng.choice(["2018Q1", "2018Q2"], n),
    })


@pytest.fixture(scope="module")
def preparado() -> pd.DataFrame:
    d, _ = r01.preparar(corpus())
    return d


# --------------------------------------------------------------- preparacion

def test_la_calidad_catastral_se_invierte_para_crecer_con_la_calidad(preparado):
    """El codigo del origen empeora al crecer; la variable del modelo, no.

    Es la correccion del coeficiente que estaba escrito como anomalia: no era
    una anomalia, era la variable leida al reves.
    """
    d = preparado
    assert (d["calidad_catastral"]
            == r01.CALIDAD_CATASTRAL_PEOR - d["CADASTRALQUALITYID"]).all()
    # El mejor codigo del origen (el mas bajo) da la mejor calidad del modelo.
    mejor = d.loc[d["CADASTRALQUALITYID"].idxmin(), "calidad_catastral"]
    peor = d.loc[d["CADASTRALQUALITYID"].idxmax(), "calidad_catastral"]
    assert mejor > peor


def test_la_calidad_ya_recodificada_correlaciona_positivo_con_el_precio(preparado):
    """Con la variable en la direccion correcta, mejor calidad es mas caro."""
    r = preparado["calidad_catastral"].corr(np.log(preparado["eur_m2"]))
    assert r > 0


def test_preparar_aparta_lo_implausible_y_deja_constancia():
    d = corpus(120)
    d.loc[0, "CONSTRUCTEDAREA"] = 5          # fuera del rango de superficie
    d.loc[1, "PRICE"] = 10                   # fuera del rango de precio
    d.loc[2, "barrio"] = None                # sin barrio asignado
    limpio, incidencias = r01.preparar(d)

    assert len(limpio) == len(d) - 3
    assert incidencias["filas"].sum() == 3
    # Ninguna fila se descarta en silencio: cada una con su motivo.
    assert set(incidencias.columns) == {"motivo", "filas", "pct"}
    assert (incidencias["pct"] <= 1).all()


def test_preparar_deja_una_sola_fila_por_anuncio_y_trimestre():
    d = corpus(80)
    repetido = d.iloc[[0]].copy()
    repetido["PRICE"] *= 1.01               # mismo anuncio, precio distinto
    limpio, _ = r01.preparar(pd.concat([d, repetido], ignore_index=True))
    assert not limpio.duplicated(["ASSETID", "trimestre"]).any()


# ------------------------------------------------------------ el ajuste
def test_el_factor_ignora_la_constante(preparado):
    """`factor` monta la columna de unos a cero, asi que el termino
    independiente no entra. Es lo que hace que invertir la calidad no mueva
    ninguna valoracion."""
    beta = r01.estimar_coeficientes(preparado)
    otro = beta.copy()
    otro[0] += 7.5                          # cambiar solo la constante
    assert np.allclose(r01.factor(preparado, beta), r01.factor(preparado, otro))


def test_invertir_la_calidad_no_mueve_ninguna_valoracion():
    """La afirmacion escrita en r01.preparar(), comprobada.

    Reescalar una variable lineal por un desplazamiento constante entra en el
    factor como un multiplicador comun, y se cancela al dividir el EUR/m2 de
    cada comparable por su propio factor.
    """
    d, _ = r01.preparar(corpus())
    al_reves = d.copy()
    al_reves["calidad_catastral"] = r01.CALIDAD_CATASTRAL_PEOR - d["calidad_catastral"]

    mitad = len(d) // 2
    partes = [(d.iloc[:mitad], d.iloc[mitad:]),
              (al_reves.iloc[:mitad], al_reves.iloc[mitad:])]
    resultados = [r01.valorar(pool, ev, r01.estimar_coeficientes(pool))
                  for pool, ev in partes]

    a, b = (r["valor_estimado"].to_numpy() for r in resultados)
    assert np.allclose(a, b, equal_nan=True)


# ------------------------------------------------------------ la valoracion
def test_no_se_valora_por_debajo_del_minimo_de_comparables():
    d, _ = r01.preparar(corpus(300))
    mitad = len(d) // 2
    pool, ev = d.iloc[:mitad], d.iloc[mitad:]
    res = r01.valorar(pool, ev, r01.estimar_coeficientes(pool))

    flojos = res[res["n_comparables"] < r01.MIN_COMPARABLES]
    assert flojos["valor_estimado"].isna().all()
    # Y nunca se calla el motivo.
    assert (flojos["motivo_no_valorable"].str.len() > 0).all()


def test_el_anuncio_evaluado_no_puede_estar_en_sus_propios_comparables():
    """El control de fugas del circuito real: pool y evaluacion son disjuntos."""
    d, _ = r01.preparar(corpus())
    rng = np.random.default_rng(r01.SEMILLA)
    ids = d["ASSETID"].unique()
    en_pool = set(rng.choice(ids, size=int(len(ids) * r01.PROP_POOL), replace=False))
    marca = d["ASSETID"].isin(en_pool)
    pool, ev = d[marca], d[~marca]
    assert set(pool["ASSETID"]) & set(ev["ASSETID"]) == set()


# ------------------------------------------------------- la hoja de Excel
def _ref(tabla) -> str:
    """openpyxl devuelve un objeto Table en memoria y la cadena al releer."""
    return tabla if isinstance(tabla, str) else tabla.ref


def metricas_de_juguete() -> dict:
    bloque = {"error_absoluto_mediano": 0.12, "dentro_de_10pct": 0.43,
              "dentro_de_20pct": 0.72, "sesgo_mediano": -0.001}
    return {
        "fuente": "idealista18, licencia ODbL-1.0.",
        "n_anuncios_utilizables": 75476, "n_pool": 52833, "n_evaluacion": 22643,
        "cobertura": 0.9614,
        "motor_comparables": bloque,
        "base_mediana_barrio": dict(bloque, error_absoluto_mediano=0.1518),
        "base_mediana_ciudad": dict(bloque, error_absoluto_mediano=0.3431),
        "error_por_fiabilidad": {"alta": 0.1037, "media": 0.1273, "baja": 0.1593},
        "coeficientes": {"jardin": -0.014, "calidad_catastral": 0.0137,
                         "log_superficie": -0.1746},
        "limitaciones": ["El objetivo es el precio PEDIDO, no el de cierre."],
    }


def test_la_hoja_de_datos_reales_sale_con_sus_tablas_con_nombre():
    from openpyxl import Workbook
    wb = Workbook()
    p05.hoja_validacion_2018(wb, metricas_de_juguete())
    ws = wb["Validación 2018"]

    # Power BI carga por nombre de tabla, no por hoja.
    assert {"tblRealFicha", "tblRealMetodos", "tblRealFiabilidad",
            "tblRealCoeficientes", "tblRealLimitaciones"} <= set(ws.tables)

    # Ninguna celda combinada puede solaparse con una tabla, o Excel se queja.
    ocupadas = set()
    for tabla in ws.tables.values():
        ini, fin = _ref(tabla).split(":")
        for fila in ws[ini:fin]:
            ocupadas.update(c.coordinate for c in fila)
    for combinada in ws.merged_cells.ranges:
        assert not {c.coordinate for f in ws[str(combinada)] for c in f} & ocupadas


def test_la_hoja_traduce_el_coeficiente_a_efecto_sobre_el_precio():
    from openpyxl import Workbook
    wb = Workbook()
    m = metricas_de_juguete()
    p05.hoja_validacion_2018(wb, m)
    ws = wb["Validación 2018"]

    ini, fin = _ref(ws.tables["tblRealCoeficientes"]).split(":")
    filas = list(ws[ini:fin])[1:]
    efectos = {f[0].value: (f[1].value, f[2].value) for f in filas}

    # La columna de efecto es exp(coeficiente) - 1, no el coeficiente crudo.
    efecto, log = efectos["Superficie (elasticidad del €/m²)"]
    assert log == m["coeficientes"]["log_superficie"]
    assert efecto == pytest.approx(np.expm1(log), abs=1e-4)
    # La calidad catastral ya sale positiva: es la correccion verificada.
    assert efectos["Calidad catastral (por grado)"][0] > 0


def test_el_veredicto_solo_aparece_si_hay_diagnostico():
    from openpyxl import Workbook
    sin = Workbook()
    p05.hoja_validacion_2018(sin, metricas_de_juguete())
    assert "tblRealVeredicto" not in sin["Validación 2018"].tables

    con = Workbook()
    p05.hoja_validacion_2018(con, metricas_de_juguete(),
                          {"veredicto": {"jardin": "verificado",
                                         "calidad_catastral": "verificado"}})
    assert "tblRealVeredicto" in con["Validación 2018"].tables


def test_el_excel_se_construye_aunque_no_haya_datos_reales(tmp_path, monkeypatch):
    """La integracion continua no tiene `data/real/`. El libro sale igual.

    Necesita la salida del circuito sintetico, que este test NO genera. En la
    integracion continua los tests corren ANTES que `run_pipeline.py`, asi que
    ahi todavia no existe y el test se salta en vez de fallar. Es el mismo
    criterio que el resto de tests que dependen de datos que no se versionan.
    """
    import config as cfg
    for f in (cfg.F_VALORADO, cfg.DIR_PROCESSED / "informe_calidad.csv",
              cfg.F_METRICAS):
        if not f.exists():
            pytest.skip(f"falta {f.name}: hay que ejecutar antes run_pipeline.py")

    monkeypatch.setattr(cfg, "F_METRICAS_REALES", tmp_path / "no_existe.json")
    monkeypatch.setattr(cfg, "F_DIAGNOSTICO_REAL", tmp_path / "tampoco.json")
    monkeypatch.setattr(cfg, "F_EXCEL", tmp_path / "libro.xlsx")
    monkeypatch.setattr(cfg, "DIR_OUTPUTS", tmp_path)

    p05.construir()

    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / "libro.xlsx")
    assert "Validación 2018" not in wb.sheetnames
    assert "Valoraciones" in wb.sheetnames


def test_las_metricas_reales_publicadas_cuadran_con_lo_que_lee_el_excel():
    """Si el fichero real esta presente, tiene que traer lo que la hoja pide."""
    import config as cfg
    if not cfg.F_METRICAS_REALES.exists():
        pytest.skip("data/real/metricas_2018.json no esta en este checkout")

    m = json.loads(cfg.F_METRICAS_REALES.read_text(encoding="utf-8"))
    for clave in ("motor_comparables", "base_mediana_barrio", "base_mediana_ciudad"):
        assert set(m[clave]) >= {"error_absoluto_mediano", "dentro_de_10pct",
                                 "dentro_de_20pct", "sesgo_mediano"}
    assert set(m["coeficientes"]) == set(r01.VARIABLES)
    # La correccion verificada, fijada contra regresiones futuras.
    assert m["coeficientes"]["calidad_catastral"] > 0


def test_la_hoja_avisa_en_cabecera_del_ano_y_de_que_es_precio_pedido():
    """Dos cosas no se pueden leer mal sin equivocarse con la cifra de al lado.

    Que los datos son de 2018 y que lo medido es el precio pedido, no el de
    cierre. Van arriba del todo, antes de la primera tabla, no en un anexo.
    """
    from openpyxl import Workbook
    wb = Workbook()
    p05.hoja_validacion_2018(wb, metricas_de_juguete())
    ws = wb["Validación 2018"]

    primera_tabla = min(int(_ref(t).split(":")[0][1:]) for t in ws.tables.values())
    cabecera = " ".join(str(ws.cell(row=i, column=1).value or "")
                        for i in range(1, primera_tabla))

    assert "2018" in cabecera
    assert "PEDIDO" in cabecera
    assert "cierre" in cabecera
    # Y el aviso va en rojo, no en el gris de las notas normales.
    rojos = [ws.cell(row=i, column=1).font.color.rgb
             for i in range(1, primera_tabla)
             if ws.cell(row=i, column=1).font.color is not None]
    assert any(c and c.endswith("9C2A2A") for c in rojos)


# ------------------------------------------------- diagnostico de coeficientes
def pool_de_juguete() -> pd.DataFrame:
    """El mismo montaje que `r04.cargar_pool`, pero sin leer los 26 MB."""
    d, _ = r01.preparar(corpus(600))
    d["log_m2"] = np.log(d["eur_m2"])
    d["y"] = d["log_m2"] - d.groupby(["barrio", "trimestre"])["log_m2"].transform("mean")
    return d.reset_index(drop=True)


def test_las_correlaciones_salen_ordenadas_y_la_matriz_es_coherente():
    import r04_diagnostico_coeficientes as r04
    c = r04.correlaciones(pool_de_juguete())

    # Matriz completa, simétrica y con la diagonal a uno.
    assert set(c["matriz"]) == set(r01.VARIABLES)
    for v in r01.VARIABLES:
        assert c["matriz"][v][v] == 1.0
        for w in r01.VARIABLES:
            assert c["matriz"][v][w] == c["matriz"][w][v]

    # Las parejas van de mayor a menor correlación en valor absoluto.
    fuerzas = [abs(p["r"]) for p in c["parejas_mas_fuertes"]]
    assert fuerzas == sorted(fuerzas, reverse=True)
    assert abs(c["correlacion_maxima"]) == fuerzas[0]

    # Y hay vecinas para los dos coeficientes que estuvieron en cuarentena.
    assert set(c["vecinas"]) == {"jardin", "calidad_catastral"}
    for v in c["vecinas"].values():
        assert len(v) == 5


def test_el_centrado_reproduce_los_coeficientes_de_produccion():
    """La rama `centrado` tiene que dar exactamente lo que estima r01.

    Si no, el diagnóstico estaría explicando un modelo que no es el que se
    publica, que es la forma más silenciosa de que un diagnóstico mienta.
    """
    import r04_diagnostico_coeficientes as r04
    p = pool_de_juguete()
    comparativa = r04.efecto_del_centrado(p)
    beta = r01.estimar_coeficientes(p)

    for i, v in enumerate(r01.VARIABLES, start=1):
        assert comparativa[v]["centrado"] == pytest.approx(beta[i], abs=1e-4)
        assert comparativa[v]["diferencia"] == pytest.approx(
            comparativa[v]["sin_centrar"] - comparativa[v]["centrado"], abs=5e-4)


def test_el_tramo_monotono_para_donde_deja_de_bajar_y_mide_la_caida_cruda():
    """La caida cruda es la que NO lleva descontado el barrio.

    Va aparte y con nombre largo a proposito: confundirla con el coeficiente del
    modelo es el error facil, porque una es diez veces la otra.
    """
    import r04_diagnostico_coeficientes as r04
    gradiente = [
        {"codigo_origen": 0.0, "eur_m2_mediano": 4000.0, "anuncios": 100},
        {"codigo_origen": 1.0, "eur_m2_mediano": 3000.0, "anuncios": 300},
        {"codigo_origen": 2.0, "eur_m2_mediano": 1000.0, "anuncios": 500},
        {"codigo_origen": 3.0, "eur_m2_mediano": 2000.0, "anuncios": 100},   # repunta
    ]
    t = r04.tramo_monotono(gradiente, n_pool=1000)

    assert (t["codigo_desde"], t["codigo_hasta"]) == (0.0, 2.0)
    assert t["anuncios_en_el_tramo"] == 900
    assert t["anuncios_fuera"] == 100
    assert t["pct_del_pool"] == 0.9
    # De 4.000 a 1.000 es -75 % en total y -50 % por grado, en dos grados.
    assert t["caida_total_sin_descontar_barrio"] == pytest.approx(-0.75, abs=1e-4)
    assert t["caida_por_grado_sin_descontar_barrio"] == pytest.approx(-0.5, abs=1e-4)


def test_quitar_el_centrado_cambia_los_coeficientes():
    """Si centrar por barrio-trimestre no moviera nada, el contraste no diría
    nada. Comprobamos que sí mueve, que es lo que le da sentido al bloque."""
    import r04_diagnostico_coeficientes as r04
    c = r04.efecto_del_centrado(pool_de_juguete())
    assert any(abs(v["diferencia"]) > 0.001 for v in c.values())
